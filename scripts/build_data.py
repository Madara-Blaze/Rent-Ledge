"""Turn the SHILPMUN workbook into seat-map JSON + a compressed header photo."""
import base64, io, json, re, sys, collections
import openpyxl
from PIL import Image

XLSX, PHOTO, OUT_JSON, OUT_IMG = sys.argv[1:5]
wb = openpyxl.load_workbook(XLSX, data_only=True)

SCHOOL_FIX = {
    "bishop cotton boy's school": "Bishop Cotton Boys' School",
    "bishop cotton boys' school": "Bishop Cotton Boys' School",
    "bishop cotton boys’ school": "Bishop Cotton Boys' School",
    "bishop cotton girls' school": "Bishop Cotton Girls' School",
    "bishop cotton girls’ school": "Bishop Cotton Girls' School",
    "devin academy for learning": "Devin Academy for Learning",
    "jain heritage school": "Jain Heritage School",
    "national public school kalkare": "National Public School, Kalkere",
    "nps kalkere": "National Public School, Kalkere",
    "national public school, kalkere": "National Public School, Kalkere",
    "vidyashilp school banashankari": "Vidyashilp School, Banashankari",
    "viswa vidyapeeth": "Vishwa Vidyapeeth",
    "chrysalis high school yehlanka": "Chrysalis High School, Yelahanka",
    "ebenezer international school bangalore": "Ebenezer International School, Bangalore",
    "rv pu north": "RV PU North",
    "vidyanikethan": "Vidyaniketan Public School",
}
COMMITTEE_FIX = {"IPL": "IPLA"}


def school(s):
    s = " ".join(str(s or "").split())
    return SCHOOL_FIX.get(s.lower(), s)


def clean(s):
    return " ".join(str(s or "").split())


# --- layout from Seat Map sheet
ROWS = []
for r in wb["Seat Map"].iter_rows(values_only=True):
    if r[0] and str(r[0]).startswith("Row "):
        row = str(r[0])[4:].strip()
        seats = [(i, c.strip()) for i, c in enumerate(r) if c and c != "|" and "ENTRY" not in c and i > 0]
        ROWS.append({"row": row, "seats": [{"id": c, "x": i} for i, c in seats]})

# --- people
people = {}
for r in wb["Seating – Final"].iter_rows(min_row=2, values_only=True):
    if not any(r[:7]):
        continue
    name, sch, reg, com, port, seat = (clean(v) for v in r[:6])
    reg_l = reg.lower()
    if "indiv" in reg_l:
        kind, reg_out = "Individual", "Individual registration"
    elif "delegation" in reg_l:
        kind, reg_out = ("Host" if school(sch) == "Vidyashilp Academy" else "Delegation"), "School delegation"
    else:
        kind, reg_out = ("Host" if school(sch) == "Vidyashilp Academy" else "Delegation"), "Class " + reg
    people[seat] = {"n": name, "s": school(sch), "c": COMMITTEE_FIX.get(com, com) or "", "p": port,
                    "r": reg_out, "k": kind}
    # Vidyashilp students only: grade number from "7-B" / "IG10-A" style class codes
    gm = re.match(r"^(?:IG)?(\d{1,2})\b", reg) if kind == "Host" else None
    if gm:
        people[seat]["g"] = int(gm.group(1))

com = None
for r in wb["EB "].iter_rows(min_row=2, values_only=True):
    if not r[1]:
        continue
    if r[0]:
        com = clean(r[0])
    seat = clean(r[3])
    assert seat not in people, seat
    people[seat] = {"n": clean(r[1]), "s": "Executive Board", "c": COMMITTEE_FIX.get(com, com),
                    "p": clean(r[2]), "r": "Executive Board", "k": "EB"}

ids = {s["id"] for row in ROWS for s in row["seats"]}
missing = [k for k in people if k not in ids]
assert not missing, missing
empty = sorted(ids - set(people))
print("seats", len(ids), "people", len(people), "empty", empty)
print(collections.Counter(p["c"] for p in people.values()))
print(collections.Counter(p["k"] for p in people.values()))

json.dump({"rows": ROWS, "people": people}, open(OUT_JSON, "w"), ensure_ascii=False, separators=(",", ":"))

# --- header photo: crop a wide band around the seats, compress
im = Image.open(PHOTO).convert("RGB")
w, h = im.size
band = im.crop((0, int(h * 0.30), w, int(h * 0.80)))
band.thumbnail((1400, 1400))
buf = io.BytesIO()
band.save(buf, "JPEG", quality=58, optimize=True, progressive=True)
open(OUT_IMG, "w").write("data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode())
print("img bytes", len(buf.getvalue()), band.size)
